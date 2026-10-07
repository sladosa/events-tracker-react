# -*- coding: utf-8 -*-
"""_excel_izbornici.py — Tip/Podtip izbornici i izgled za app Excel koji pišu Python alati. S164.

ZAŠTO POSTOJI
    `visa_popravak.pisi()` (S148) je Podtipu davao RAVAN popis svih Podtipova — pa izbornik
    nije pratio Tip, a Podtip mimo `validation_rules` se uveze kao tekst BEZ GREŠKE
    (CLAUDE.md). Saša je to prijavio dvaput (S148 i S164). App export (`excelExport.ts`
    `addDependentDropdowns`) to radi ispravno od S105: imenovani raspon po vrijednosti Tipa
    (`Dep_tip_<Tip>`) i `INDIRECT` na Podtip stupcu. Ovaj modul je ISTO pravilo u Pythonu.

⚠ JEDINA IMPLEMENTACIJA za app-format fileove. Tko piše Tip/Podtip izbornik, zove ovo —
  ravan popis Podtipova je kvar, ne pojednostavljenje.
  (Stariji review alati iz S107 — `normalize_financije`, `sync_taxonomy`, `suggest_candidates`,
  `apply_ai` — pišu vlastite `Tip_*` izbornike za REVIEW fileove, ne za app uvoz.)

⚠ BRANA, NE DISCIPLINA: `provjeri()` ponovo otvori spremljeni file i za SVAKI Tip izračuna što
  bi formula u Excelu dala (ista SUBSTITUTE pravila) i postoji li taj imenovani raspon s
  točnim opcijama. Pisac (`visa_popravak.pisi`) staje ako provjera padne.
"""
from __future__ import annotations

import re

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# Isto kao `DIACRITICAL_MAP` u `src/lib/excelExport.ts`.
DIJAKRITICI = [('č', 'c'), ('ć', 'c'), ('š', 's'), ('ž', 'z'), ('đ', 'd'),
               ('Č', 'C'), ('Ć', 'C'), ('Š', 'S'), ('Ž', 'Z'), ('Đ', 'D')]
PREFIKS = 'Dep_tip'          # = sanitizeNamedRange(`Dep_${'tip'}`) u app exportu
LIST = 'DropdownData'


def ime_raspona(vrijednost: str) -> str:
    """= `sanitizeNamedRange` iz excelExport.ts."""
    s = f'{PREFIKS}_{vrijednost}'
    for a, b in DIJAKRITICI:
        s = s.replace(a, b)
    s = re.sub(r'[^A-Za-z0-9_]', '_', s)
    return re.sub(r'^(\d)', r'_\1', s)


def _zamjene(tipovi) -> list[tuple[str, str]]:
    """SUBSTITUTE parovi za znakove koji se STVARNO javljaju u imenima Tipova.
    Puni lanac app exporta (20 zamjena) je duži od nužnog; kraći je i lakši za provjeru."""
    znakovi = sorted({c for t in tipovi for c in t if not re.match(r'[A-Za-z0-9_]', c)})
    dij = dict(DIJAKRITICI)
    return [(c, dij.get(c, '_')) for c in znakovi]


def _formula_ime(celija: str, zamjene) -> str:
    sub = celija
    for a, b in zamjene:
        sub = f'SUBSTITUTE({sub},"{a}","{b}")'
    return f'INDIRECT("{PREFIKS}_"&{sub})'


def _excel_ime(tip: str, zamjene) -> str:
    """Što Excelova formula izračuna za danu vrijednost Tipa (simulacija SUBSTITUTE lanca)."""
    s = tip
    for a, b in zamjene:
        s = s.replace(a, b)
    return f'{PREFIKS}_{s}'


def tip_podtip_izbornici(wb, ws, col_tip: int, col_pod: int, prvi: int, zadnji: int,
                         tax: dict[str, list[str]]):
    """Tip = popis Tipova (+ `N/A`), Podtip = popis OVISAN o Tipu u istom retku."""
    dd = wb.create_sheet(LIST)
    dd.sheet_state = 'hidden'
    tipovi = sorted(tax) + ['N/A']
    for i, t in enumerate(tipovi, start=1):
        dd.cell(i, 1, t)
    stupac = 2
    for t in sorted(tax):
        opcije = tax[t]
        if not opcije:
            continue
        dd.cell(1, stupac, f'tip={t}')
        for i, p in enumerate(opcije, start=2):
            dd.cell(i, stupac, p)
        L = get_column_letter(stupac)
        wb.defined_names[ime_raspona(t)] = DefinedName(
            ime_raspona(t), attr_text=f"{LIST}!${L}$2:${L}${1 + len(opcije)}")
        stupac += 1

    T, P = get_column_letter(col_tip), get_column_letter(col_pod)
    dv_t = DataValidation(type='list', formula1=f'{LIST}!$A$1:$A${len(tipovi)}', allow_blank=True)
    ws.add_data_validation(dv_t)
    dv_t.add(f'{T}{prvi}:{T}{zadnji}')
    # Relativna adresa prvog retka: Excel je pomiče za svaki redak raspona.
    zam = _zamjene(tax)
    f = _formula_ime(f'{T}{prvi}', zam)
    dv_p = DataValidation(type='list', formula1=f, allow_blank=True, showErrorMessage=True,
                          errorTitle='Podtip', error='Podtip ne pripada odabranom Tipu.')
    ws.add_data_validation(dv_p)
    dv_p.add(f'{P}{prvi}:{P}{zadnji}')
    # Izbornik ne provjerava vrijednost koja je već upisana: promijeni li se Tip, stari Podtip
    # ostane. Crveno je jedini znak da je par sad nevaljan.
    ws.conditional_formatting.add(
        f'{P}{prvi}:{P}{zadnji}',
        FormulaRule(formula=[f'AND({P}{prvi}<>"",{T}{prvi}<>"N/A",'
                             f'ISERROR(MATCH({P}{prvi},{f},0)))'],
                    fill=PatternFill('solid', fgColor='F8696B')))
    return len(f)


def provjeri(path, tax: dict[str, list[str]]) -> list[str]:
    """Ponovo otvori file i provjeri da formula za SVAKI Tip pogađa raspon s točnim opcijama."""
    wb = load_workbook(path)
    greske = []
    zam = _zamjene(tax)
    imena = wb.defined_names
    for t, opcije in tax.items():
        if not opcije:
            continue
        ime = _excel_ime(t, zam)
        if ime != ime_raspona(t):
            greske.append(f'Tip „{t}": formula daje {ime}, raspon se zove {ime_raspona(t)}')
            continue
        if ime not in imena:
            greske.append(f'Tip „{t}": nema imenovanog raspona {ime}')
            continue
        (list_, ref), = list(imena[ime].destinations)
        vrijednosti = [c.value for red in wb[list_][ref.replace('$', '')] for c in red]
        if vrijednosti != list(opcije):
            greske.append(f'Tip „{t}": raspon nosi {vrijednosti}, očekivano {list(opcije)}')
    pod = [dv for dv in wb['Events'].data_validations.dataValidation if 'INDIRECT' in (dv.formula1 or '')]
    if not pod:
        greske.append('Podtip stupac nema INDIRECT izbornik (ravan popis?)')
    elif len(pod[0].formula1) > 255:
        greske.append(f'Formula izbornika ima {len(pod[0].formula1)} znakova (Excel prima 255)')
    return greske


def izgled_pregleda(ws, hdr: int, heads: list[str]):
    """Sašin izgled za pregled (S164, sa `na_razvrstavanje_*` fileova): legenda sklopljena,
    tehnički stupci u sklopljenim grupama, zamrznuti datum + opis, širi opis."""
    def col(name):
        return heads.index(name) + 1

    ws.sheet_format.outlineLevelRow = 1
    ws.sheet_format.outlineLevelCol = 1
    for r in range(1, hdr - 1):                       # legenda + prazan red; naslov ostaje
        ws.row_dimensions[r].outlineLevel = 1
        ws.row_dimensions[r].hidden = True
    for grupa in (['event_id', 'Area', 'Category_Path'],
                  ['session_start', 'created_at', 'User'],
                  ['Racun', 'Izvor', 'Smjer']):
        cs = [col(n) for n in grupa if n in heads]
        for c in cs:
            d = ws.column_dimensions[get_column_letter(c)]
            d.outlineLevel, d.hidden = 1, True
        ws.column_dimensions[get_column_letter(max(cs) + 1)].collapsed = True
    ws.column_dimensions[get_column_letter(col('leaf comment'))].width = 37
    ws.freeze_panes = ws.cell(hdr + 1, col('Uplata'))
