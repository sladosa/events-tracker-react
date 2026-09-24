# -*- coding: utf-8 -*-
"""
visa_popravak.py  (S148, 2026-09-24)
====================================
Excel za uvoz kroz aplikaciju koji popravlja Visa kosare 2026. Nalaz je od
`visa_kosare.py`: svih 7 PBZ izvoda 2026. slaze se s RF naplatom U CENT, a
visak je u BAZI — Sasini rucni retci (sheet `sasa EU` u Kokinoj Excelici)
uvezeni UZ retke s izvoda, jer se dedup `(datum, iznos)` nije poklopio:

  * isti dan, isti iznos, dvaput                 (DM, Spar, Amsterdam 6,05)
  * MJESEC RANIJE od bankinog datuma             (Konzum x3, Decathlon — 06.06. vs 06.07.)
  * dan ranije                                   (Amsterdam 30,20)
  * rucna rata uz bankinu ratu                   (Konzum 2/4 x2 i 3/4, Traperice 4/4,
                                                  Tekstilpromet 2/2, Satrak 2/10)
  * rata u krivoj kosari (`Datum naplate`)       (Bauhaus 2/6, Pedikura 2/2)

/!\ POPIS, NE PRAVILO. Svaki redak je provjeren pojedinacno protiv izvoda (isti
    razred kao `RUCNO` u `rate_alat.py`). Automatsko "brisi sve bez `Izvod opis`"
    bi obrisalo i 8 redaka za koje odgovor zna samo Sasa (v. NEPOZNATO).

/!\ BRISE SE SAMO UZ ZIVOG PARA. Prije pisanja filea alat provjeri da za svaki
    redak koji brise u bazi postoji par s ISTIM iznosom i POTVRDOM izvoda
    (`Izvod opis`). Nema li ga — alat STANE: brisanje bez para gubi transakciju.

/!\ Opis ide s obrisanog retka na zadrzani (Kokin/Sasin tekst je autoritet za
    opis, izvod za iznos — S114). Inace bi `Konzum 2/4` postao samo `Konzum`.

/!\ UVOZI KOKA (vlasnica Aree i autorica svih ovih redaka u bazi — `user_id`
    izmjeren) — samo tada `Delete?` smije brisati; tudji redak parser odbije.

Pokretanje:
  $env:ET_TARGET='prod'; Financije\\run.bat visa_popravak.py            -> dry run
  $env:ET_TARGET='prod'; Financije\\run.bat visa_popravak.py --file     -> + xlsx
"""
from __future__ import annotations

import os
import sys
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, str(Path(__file__).parent))
from _db import ROOT, load_env  # noqa: E402
from rate_alat import AREA, ATTRS, CAT_PATH, KOKA_EMAIL, col_letter, tekst  # noqa: E402
from uskladi_izvod import load_db  # noqa: E402

ZG = ZoneInfo('Europe/Zagreb')

# /!\ KOLONA G MORA NOSITI E-MAIL AUTORA RETKA, NE UVOZNIKA (S148, placeno).
#     Uvoz po njoj odlucuje je li redak "tudji". Kokin e-mail na Sasinom retku
#     => parser ga vidi kao njezin, baza kaze da je Sasin, `canUpdateExisting`
#     odbije i redak ode u CREATE: 7 DUPLIKATA uz poruku "not found in database".
#     S tocnim e-mailom uvoz ponudi `fix_as_owner` i ispravi original na mjestu.
EMAIL = {
    'eeb78414-849e-425b-8825-15ef09583dfe': KOKA_EMAIL,
    '768a6056-91fd-42bb-98ae-ee83e6bd6c8d': 'sasasladoljev59@gmail.com',
}
OUT = ROOT / 'data-prep_data' / 'Financije'
EPS = 0.005

# (id-prefiks, zadrzani par, zasto)
BRISI = [
    ('c63ee548', '2635bfd9', 'DM 13,41 · 09.01. — isti dan i iznos dvaput; par nosi izvod'),
    ('9420723a', 'cf3f869b', 'Spar 24,92 · 22.01. — isti dan i iznos dvaput; par nosi izvod'),
    ('5110069f', 'e144365c', 'Amsterdam 6,05 · 26.05. — isti dan i iznos dvaput; par nosi izvod'),
    ('4f6a79ab', '203c24f9', 'Amsterdam 30,20 · 25.05. — banka 26.05. (Buno Coffee)'),
    ('45ac82b1', '4795a51f', 'Konzum 21,57 · 06.06. — banka 06.07., mjesec ranije'),
    ('5a9bd57e', '2af40bd6', 'Konzum 5,27 · 06.06. — banka 06.07., mjesec ranije'),
    ('9465bf76', '47af402e', 'Konzum 12,23 · 07.06. — banka 07.07., mjesec ranije'),
    ('2e31de15', 'ae5b2343', 'Decathlon 23,96 · 07.06. — banka 07.07., mjesec ranije'),
    ('f76fbea8', '645e6537', 'Konzum 2/4 · 10.04. — banka: RATA 02/04 kupnje 10.03.'),
    ('a5b890cb', 'f0c3f806', 'Konzum 2/4 · 11.05. — banka: RATA 03/04 kupnje 10.03.'),
    ('26f078f1', '1255714a', 'Konzum 3/4 · 11.06. — banka: RATA 04/04 kupnje 10.03.'),
    ('d7f122d4', 'd91b4369', 'Traperice 4/4 · 17.06. — banka: RATA 04/04 kupnje 17.02.'),
    ('307766a4', '5668f16c', 'Tekstilpromet 2/2 35,00 — banka: RATA 02/02 = 34,99'),
    ('e53c90f0', '5dff0591', 'Satrak 2/10 98,86 — banka: RATA 02/10 = 98,80'),
]

# (id-prefiks, {atribut: nova vrijednost}, novi komentar ili None, zasto)
ISPRAVI = [
    ('119d7f5b', {'Datum naplate': date(2026, 8, 7),
                  'Izvod opis': 'RATA 02/ 06-BAUHAUS - JANKOMIR - ZAGREB-SUSEDGRAD - ŠKORPIKO'},
     None, 'Bauhaus 2/6 je na izvodu PBZVIZA_2026-07 ⇒ naplata 07.08., ne 06.07.'),
    ('e45678c1', {'Datum naplate': date(2026, 8, 7),
                  'Izvod opis': 'RATA 02/ 02-PONUDA DANA.HR - ZAGREB - ZVEČAJSKA 17'},
     None, 'Pedikura 2/2 je na izvodu PBZVIZA_2026-07 ⇒ naplata 07.08., ne 06.07.'),
    ('d91b4369', {'Rata br': 4}, 'Traperice 4/4',
     'izvod kaze RATA 04/04, redak nosi 3/4 (a 3/4 vec postoji)'),
    ('645e6537', {'Rata br': 2}, 'Konzum 2/4', 'opis s obrisanog rucnog retka'),
    ('f0c3f806', {'Rata br': 3}, 'Konzum 3/4', 'opis s obrisanog rucnog retka'),
    ('1255714a', {'Rata br': 4}, 'Konzum 4/4', 'opis s obrisanog rucnog retka'),
    ('5668f16c', {}, 'Tekstilpromet 2/2', 'opis s obrisanog rucnog retka (bio sirovi tekst izvoda)'),
    ('5dff0591', {'Tip': 'auto C5', 'Podtip': 'popravci'}, 'Šatrak 2/10',
     'opis i Tip/Podtip s obrisanog rucnog retka (bankin nosi N/A)'),
    ('d0155b1c', {'Izvor': 'Racun', 'Racun': 'Kokin tekući ZABA', 'Status': 'Izvrsen',
                  'Datum naplate': date(2026, 9, 23)},
     None, 'PP 8,60 placen nalogom s Kokinog ZABA (Sasa, S148) — ne Visa'),
]

# Nije u fileu: odgovor zna samo Sasa. Nema ih ni na jednom PBZ/MC/ZABA izvodu
# 2025-2026; izvor im je Sasin sheet `sasa EU` (Izvor = Visa).
NEPOZNATO = ['b9abd143', 'dc8f28b6', '1603c096', 'f3cc8a22',
             'caa602b1', '6d7bb025', '4eb5059d', '88378b95']

# Trag s izvoda za svaki (pretraga +-10 dana, +-10 % iznosa). Cetiri imaju
# bankin redak ISTOG DANA sa zamijenjenom/krivom znamenkom, koji je u bazi vec
# potvrdjen => vjerojatno duplikat s tipfelerom, ne Cash.
SAVJET = {
    'b9abd143': 'Carglass 85,00 — nigdje (ni PBZ, MC, ZABA); kandidat za Cash',
    'dc8f28b6': 'Biberon 9,10 — banka ISTI DAN Biberon 9,01 (vec u bazi) ⇒ vjerojatno duplikat, obrisi',
    '1603c096': 'Bates 1/3 164,68 — banka ISTI DAN RATA 01/03 163,68 (vec u bazi) ⇒ duplikat, obrisi',
    'f3cc8a22': 'Pekara 3,60 — isti dan banka ima 3 male: 3,40 Dubravica, 2,00 Junior, 1,60 Svetice; nejasno',
    'caa602b1': "McDonalds 7,50 — banka ISTI DAN McDonald's 7,40 (vec u bazi) ⇒ vjerojatno duplikat, obrisi",
    '6d7bb025': 'H&M 49,67 — banka ISTI DAN H&M 46,97 (zamijenjene znamenke, vec u bazi) ⇒ duplikat, obrisi',
    '4eb5059d': 'Amsterdam 8,60 — nema bliskog iznosa na izvodu; kandidat za Cash',
    '88378b95': 'Cestarina 13,50 — nigdje; kandidat za Cash (ENC/gotovina?)',
}
SAVJET['f3cc8a22'] = ('Pekara 3,60 · 05.05. — banka isti dan Pekarnica Junior 2,00 + '
                      'Pekarna Svetice 1,60 = 3,60 (oboje vec u bazi) ⇒ 1:N duplikat')

# original (Sasin) -> kopija (Kokina), izmjereno nakon uvoza 24.09.2026. 19:15
KOPIJE = {
    'e4317c28': '60611fbe', 'c8181fb6': '50e3af8b', '7b44f3a4': '23e4bd6b',
    '36ea1c29': '8a408c10', 'a5a2c277': 'e0dd17b4', 'ebb614b2': '3137a6f8',
    '4485d4cf': '69ef6ea8',
}

# Sasine odluke (S148). Carglass 85,00 namjerno NIJE ovdje — jos otvoren.
ODLUKA = {
    '4eb5059d': 'Cash', '88378b95': 'Cash',
    'dc8f28b6': 'DELETE', '1603c096': 'DELETE', 'caa602b1': 'DELETE',
    '6d7bb025': 'DELETE', 'f3cc8a22': 'DELETE',
}


def num(v):
    try:
        return round(float(v or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def to_date(v):
    if v is None or v == '':
        return None
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v)[:10])


def lokalno(ss: str) -> str:
    """/!\\ Baza drzi UTC, file nosi LOKALNO `HH:MM` — krivi sat POMICE redak."""
    return datetime.fromisoformat(ss).astimezone(ZG).strftime('%H:%M')


def pick(by8, pref):
    hit = [r for k, r in by8.items() if k == pref]
    if len(hit) != 1:
        sys.exit(f'✗ id {pref}: nadjeno {len(hit)} redaka — popis vise ne odgovara bazi.')
    return hit[0]


def main():
    target = os.environ.get('ET_TARGET', 'test').strip().lower()
    url, key = load_env(target)
    print(f'[{target.upper()}] {url}')
    if target != 'prod':
        print('⚠ Popis je izmjeren na PROD-u; na TEST-u ovi id-evi ne postoje.')
    rows = load_db(url, key)
    by8 = {r['id'][:8]: r for r in rows}

    if '--duplikati' in sys.argv:
        # S148: prvi uvoz kolovoza napravio je kopije 7 Sasinih redaka (v. EMAIL).
        # Kopija (Kokina, nosi ispravne vrijednosti) se brise; njene vrijednosti
        # idu na ORIGINAL, koji zadrzava autorstvo.
        redci = []
        for orig, kop in KOPIJE.items():
            o, k = pick(by8, orig), pick(by8, kop)
            if o['user_id'] == k['user_id'] or o['event_date'] != k['event_date']:
                sys.exit(f'✗ {orig}/{kop} nisu par original/kopija')
            redci.append((k, dict(k['attrs']), k.get('comment'), True,
                          f'OBRIŠI — kopija Sasinog retka {orig} (uvoz S148, kriv e-mail)'))
            redci.append((o, dict(k['attrs']), k.get('comment'), False,
                          f'ISPRAVI — vrijednosti s kopije {kop}; autor ostaje Sasa'))
        return pisi(redci, by8, 'visa_duplikati')

    if '--odluka' in sys.argv:
        # Treci krug: Sasine odluke nad OTVORENIM (S148). `Cash` zadrzava racun,
        # a `Datum naplate` = dan troska (konvencija svih postojecih Cash redaka).
        redci = []
        for d, odl in ODLUKA.items():
            r = pick(by8, d)
            a = dict(r['attrs'])
            if odl == 'DELETE':
                redci.append((r, a, r.get('comment'), True, 'OBRIŠI — ' + SAVJET[d]))
            else:
                a.update({'Izvor': 'Cash', 'Datum naplate': to_date(r['event_date'])})
                redci.append((r, a, r.get('comment'), False, 'CASH — Sasina odluka S148'))
        return pisi(redci, by8, 'visa_odluka')

    if '--otvoreno' in sys.argv:
        # Drugi krug (nakon uvoza prvog filea): 8 redaka BEZ izmjene, Sasa sam
        # mijenja `Izvor` (npr. Cash) i uvozi. Savjet je trag s izvoda, ne odluka.
        redci = [(pick(by8, d), dict(pick(by8, d)['attrs']), pick(by8, d).get('comment'),
                  False, 'OTVORENO — ' + SAVJET.get(d, '')) for d in NEPOZNATO]
        return pisi(redci, by8, 'visa_otvoreno')

    # -- invarijanta: svaki brisani ima zivog, potvrdjenog para istog iznosa
    greske = 0
    for d, par, _ in BRISI:
        r, p = pick(by8, d), pick(by8, par)
        ok = (p['attrs'].get('Izvod opis') and p['attrs'].get('Izvor') == 'Visa'
              and (abs(num(r['attrs'].get('Isplata')) - num(p['attrs'].get('Isplata'))) < EPS
                   or d in ('307766a4', 'e53c90f0')))   # rata: banka 34,99 / 98,80
        if not ok:
            greske += 1
            print(f'✗ {d}: par {par} nije potvrdjen/istog iznosa')
    if greske:
        sys.exit('✗ Stajem — brisanje bez para gubi transakciju.')
    print(f'✓ {len(BRISI)} brisanja, svako uz zivog potvrdjenog para')

    s_del = sum(num(pick(by8, d)['attrs'].get('Isplata')) for d, _, _ in BRISI)
    s_unk = sum(num(pick(by8, d)['attrs'].get('Isplata')) for d in NEPOZNATO)
    print(f'  Σ brisanja {s_del:.2f} · 2 rate u pravu kosaru · Σ nepoznatih {s_unk:.2f}')
    print('\nNEPOZNATO (ceka Sasu):')
    for d in NEPOZNATO:
        r = pick(by8, d)
        print(f'  {r["event_date"]}  {num(r["attrs"].get("Isplata")):8.2f}  {r.get("comment")}')

    if '--file' not in sys.argv:
        print('\n(dry run — `--file` pise xlsx)')
        return

    redci = []   # (redak, attrs, komentar, delete?, zasto)
    for d, par, zasto in BRISI:
        r = pick(by8, d)
        redci.append((r, dict(r['attrs']), r.get('comment'), True, 'OBRIŠI — ' + zasto
                      + f' (ostaje {par})'))
    for d, nove, kom, zasto in ISPRAVI:
        r = pick(by8, d)
        a = dict(r['attrs'])
        a.update(nove)
        redci.append((r, a, kom if kom is not None else r.get('comment'), False,
                      'ISPRAVI — ' + zasto))
    pisi(redci, by8, 'visa_popravak')


def pisi(redci, by8, ime, tax=None):
    """`tax` = {Tip: [Podtip...]} -> dropdowni Tip/Podtip na svakom retku.
    /!\ Podtip mimo `validation_rules` uveze se kao tekst BEZ GRESKE (CLAUDE.md),
        pa file koji covjek rucno klasificira mora nuditi samo valjane vrijednosti."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Events'
    purple = PatternFill('solid', fgColor='7030A0')
    blue = PatternFill('solid', fgColor='4472C4')
    yellow = PatternFill('solid', fgColor='FFF2CC')
    red = PatternFill('solid', fgColor='F8CBAD')
    white = Font(bold=True, color='FFFFFF')

    ws.cell(1, 1, 'ATTRIBUTE LEGEND:').font = Font(bold=True, size=12)
    for i, h in enumerate(['Col', 'Area', 'Category_Path', 'Attribute', 'Type', 'Unit'], start=1):
        c = ws.cell(2, i, h)
        c.fill, c.font, c.alignment = purple, white, Alignment(horizontal='center')
    for i, (name, typ) in enumerate(ATTRS):
        r = 3 + i
        ws.cell(r, 1, col_letter(9 + i))
        ws.cell(r, 2, AREA)
        ws.cell(r, 3, CAT_PATH)
        ws.cell(r, 4, name)
        ws.cell(r, 5, typ)
    top = 3 + len(ATTRS) + 1
    ws.cell(top, 1, 'EVENT DATA:').font = Font(bold=True, size=12)
    hdr = top + 1
    del_col = 9 + len(ATTRS)
    heads = (['event_id', 'Area', 'Category_Path', 'event_date', 'session_start',
              'created_at', 'User', 'leaf comment'] + [n for n, _ in ATTRS] + ['Delete?'])
    for i, h in enumerate(heads, start=1):
        c = ws.cell(hdr, i, h)
        c.fill, c.font = blue, white

    for j, (r, a, kom, dele, _, *extra) in enumerate(redci):
        # r = None je NOVI redak: bez event_id, datum i minuta dolaze u `extra`
        rr = hdr + 1 + j
        stari = r['attrs'] if r else {}
        tekst(ws, rr, 1, r['id'] if r else None)
        tekst(ws, rr, 2, AREA)
        tekst(ws, rr, 3, CAT_PATH)
        tekst(ws, rr, 4, r['event_date'] if r else extra[0][0])
        c = ws.cell(rr, 5, lokalno(r['session_start']) if r else extra[0][1])
        c.number_format = '@'
        if r and r['user_id'] not in EMAIL:
            sys.exit(f'✗ nepoznat autor {r["user_id"]} — dopuni EMAIL')
        tekst(ws, rr, 7, EMAIL[r['user_id']] if r else KOKA_EMAIL)
        tekst(ws, rr, 8, kom or None)
        if kom != (r.get('comment') if r else None):
            ws.cell(rr, 8).fill = yellow
        for i, (name, typ) in enumerate(ATTRS):
            v = a.get(name)
            col = 9 + i
            if v is None or v == '':
                continue
            if typ == 'boolean':
                tekst(ws, rr, col, 'true' if v in (True, 'true') else 'false')
            elif typ == 'datetime':
                dd = to_date(v)
                ws.cell(rr, col, datetime(dd.year, dd.month, dd.day, 12, 0)).number_format = 'dd.mm.yyyy'
            elif typ == 'number':
                ws.cell(rr, col, float(v))
            else:
                tekst(ws, rr, col, v)
            if str(v) != str(stari.get(name)):
                ws.cell(rr, col).fill = yellow
        if dele:
            tekst(ws, rr, del_col, 'DELETE')
            for cc in range(1, del_col + 1):
                if ws.cell(rr, cc).fill != yellow:
                    ws.cell(rr, cc).fill = red
    ws.auto_filter.ref = f'A{hdr}:{col_letter(del_col)}{hdr + len(redci)}'
    if tax:
        from openpyxl.worksheet.datavalidation import DataValidation
        orange = PatternFill('solid', fgColor='F4B084')
        dd = wb.create_sheet('DropdownData')
        podtipovi = sorted({p for v in tax.values() for p in v})
        for i, t in enumerate(sorted(tax), start=1):
            dd.cell(i, 1, t)
        for i, p in enumerate(podtipovi, start=1):
            dd.cell(i, 2, p)
        ctip = 9 + [n for n, _ in ATTRS].index('Tip')
        cpod = ctip + 1
        prvi, zadnji = hdr + 1, hdr + len(redci)
        for col, n in ((ctip, len(tax)), (cpod, len(podtipovi))):
            L = 'A' if col == ctip else 'B'
            dv = DataValidation(type='list', formula1=f'DropdownData!${L}$1:${L}${n}',
                                allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f'{col_letter(col)}{prvi}:{col_letter(col)}{zadnji}')
        for rr in range(prvi, zadnji + 1):
            if ws.cell(rr, ctip).value in (None, 'N/A'):
                ws.cell(rr, ctip).fill = orange
                ws.cell(rr, cpod).fill = orange
        # parovi, da se vidi koji Podtip pripada kojem Tipu
        tp = wb.create_sheet('Tip-Podtip')
        tp.append(['Tip', 'Podtip'])
        for t in sorted(tax):
            for p in tax[t]:
                tp.append([t, p])
        tp.column_dimensions['A'].width = 14
        tp.column_dimensions['B'].width = 40
    for i, w in enumerate([38, 14, 12, 11, 8, 8, 12, 22], start=1):
        ws.column_dimensions[col_letter(i)].width = w
    for i in range(len(ATTRS) + 1):
        ws.column_dimensions[col_letter(9 + i)].width = 14

    pr = wb.create_sheet('Pregled')
    pr.append(['event_date', 'iznos', 'opis', 'što i zašto'])
    for r, a, kom, _, zasto, *extra in redci:
        pr.append([r['event_date'] if r else extra[0][0], num(a.get('Isplata')), kom, zasto])
    if ime in ('visa_popravak', 'visa_otvoreno'):
        pr.append([])
        pr.append(['NIJE U FILEU — odgovor zna samo Saša (nema ih ni na PBZ, MC ni ZABA izvodu):'])
        for d in NEPOZNATO:
            r = pick(by8, d)
            pr.append([r['event_date'], num(r['attrs'].get('Isplata')), r.get('comment'), r['id']])
    for col, w in zip('ABCD', (12, 10, 22, 90)):
        pr.column_dimensions[col].width = w

    out = OUT / f'{ime}_{datetime.now():%Y%m%d_%H%M}.xlsx'
    wb.save(out)
    print(f'\n→ {out}  ({len(redci)} redaka)')


if __name__ == '__main__':
    main()
