# -*- coding: utf-8 -*-
"""razvrstaj_na.py — prijedlog `Tip`/`Podtip` za `N/A` retke, kao app Excel za uvoz. S164.

ZASTO POSTOJI
    Pločica „Kamo ide novac" (`docs/RAZREZ_SPEC.md` §4.5) prikazuje N/A kao vlastitu
    krišku. U 10/2025–09/2026 to je 121 redak / 2.110,80 — 79 iz 10–12/2025, prije nego
    je Koka prešla na unos u appu (Tip obavezan). 112 ih nosi `Izvod opis`, pa ih
    povijest zna razvrstati.

KAKO PREDLAŽE — brojanjem povijesti, nikad pogađanjem (pravila: data-prep_tools/CLAUDE.md
„Rječnik `Izvod opis`")
    0. OPIS: kratka ljudska oznaka u komentaru (≤ 30 znakova, nije prepisan tekst izvoda)
       — „Kokin opis je jači od statistike"
    1. KARTICA: trgovac (`uvezi_transu.kljuc`: bez `[kartica: …]`, rate, broja transakcije)
       2. trgovac + iznos, kad trgovac sam nije jednoglasan (Apple 2,99 = Cloud backup 17/17)
    3. RAČUN: `presedani.Presedani` — primatelj + poziv na broj → primatelj → iznos s predznakom
    Svaki ključ traži ≥ 90 % jednoglasnosti i ≥ 3 presedana (ime+poziv: ≥ 2). `N/A` u
    povijesti NE glasa. Posrednik (KEKS PAY, KUPOVINA…, goli PAYPAL) se ne pokušava.
    Prijedlog mimo `validation_rules` se odbacuje (uvezao bi se kao tekst bez greške).
    Što ne prođe — ostaje `N/A` u fileu (narančasto) za ručno razvrstavanje.

IZLAZ (`--file`) — app Excel, jedan po RAČUNU (tko pregledava): `na_razvrstavanje_<račun>_*.xlsx`
    - kol. G = autor retka (S148), Tip/Podtip padajući izbornici, prijedlog žuto,
      bez prijedloga narančasto; list `Pregled` nosi dokaz za svaki redak
    - Podtip izbornik OVISI o Tipu (`_excel_izbornici`, S164); izgled = Sašin (legenda i
      tehnički stupci sklopljeni, zamrznuti datum + opis)
    - uvozi VLASNICA Aree (Koka ili Saša pod njezinim računom); tuđi retci: „fix as owner"
    - pregled uvoza smije pokazati SAMO izmjene (`N New = 0`, `Delete = 0`)

PROVJERA UNATRAG (`--provjera`)
    Već razvrstane retke u prozoru glumi kao N/A, s poviješću STRIKTNO STARIJOM od retka
    („bi li alat pogodio u trenutku kad je redak nastao"), i broji pogotke. Mjeri koliko
    vrijede prijedlozi prije nego ih itko pregledava.

Pokretanje (PowerShell, iz data-prep_tools\\):
    Financije\\run.bat razvrstaj_na.py                      zadnjih 12 punih mjeseci, ispis
    Financije\\run.bat razvrstaj_na.py --file               + Excel po računu
    Financije\\run.bat razvrstaj_na.py --od 2025-01-01 --do 2025-12-31 --file
    Financije\\run.bat razvrstaj_na.py --sve --file         cijela povijest (1.513 redaka)
    Financije\\run.bat razvrstaj_na.py --provjera           točnost na razvrstanim retcima
    Financije\\run.bat razvrstaj_na.py --file --preuzmi <razvrstan.xlsx> [<drugi.xlsx>]
                    ručne odluke iz ranijeg filea (npr. TEST) po event_id u svjež file OVE baze.
                    TEST file se na PROD NE uvozi (kol. G = TEST vlasnik, PROD retke je upisala Koka).
    Baza: $env:ET_TARGET='prod' (bez toga TEST).
"""
from __future__ import annotations

import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
sys.path.insert(0, str(Path(__file__).parent))
from _db import load_env, rest, target  # noqa: E402
from presedani import MAX_OZNAKA, MIN_PO_IMENU, MIN_UDIO, RATA_RE, Presedani  # noqa: E402
from uskladi_izvod import load_db, net  # noqa: E402
from uvezi_transu import RUCNO  # noqa: E402
from uvezi_transu import kljuc as _kljuc_mc  # noqa: E402
from visa_uvoz_izvoda import KLASA, taksonomija  # noqa: E402
import visa_popravak  # noqa: E402
from _excel_izbornici import izgled_pregleda  # noqa: E402

KARTICE = ('Mastercard', 'Visa')
# Posrednik nije trgovac: `Izvod opis` ne govori što je kupljeno (CLAUDE.md S124).
POSREDNIK = re.compile(r'^(keks pay|kupovina|paypal\s*$|paypal\s*\*\s*$)', re.I)


def arg(name, default=None):
    if name in sys.argv:
        i = sys.argv.index(name)
        return sys.argv[i + 1] if i + 1 < len(sys.argv) else default
    return default


def prozor():
    if '--sve' in sys.argv:
        return '0000-01-01', '9999-12-31'
    d = date.today().replace(day=1)
    do_ = (d.replace(day=1) - __import__('datetime').timedelta(days=1)).isoformat()
    od_ = d.replace(year=d.year - 1).isoformat()
    return arg('--od', od_), arg('--do', do_)


def kljuc_kartice(opis) -> str:
    """Trgovac. Visa izvod nosi `TRGOVAC - ADRESA - GRAD`, MC samo trgovca: bez rezanja
    na ` - ` `INA BP - KSAVERSKA CESTA 2F - ZAGREB` nikad ne nađe `INA BP` (isto pravilo
    kao `visa_uvoz_izvoda.trgovac`)."""
    return _kljuc_mc(str(opis or '').split(' - ')[0])


def rucno(k, iznos):
    """Ručni rječnici koji već postoje (S124 `uvezi_transu`, S148 `visa_uvoz_izvoda`)."""
    # /!\ `RUCNO_IZNOS` (Apple 9,99 -> HBOmax) se NAMJERNO ne koristi: bila je odluka za
    #     jednu transu, a provjera unatrag ju je 3x pobila (Cloud backup na istih 9,99).
    for rk, par in RUCNO.items():
        if k.startswith(rk):
            return par, f'ručni rječnik „{rk}"'
    for rk, par in KLASA.items():
        if rk.lower() in k:
            return par, f'ručni rječnik „{rk}"'
    return None, None


def fold(s) -> str:
    s = unicodedata.normalize('NFKD', str(s or ''))
    return ''.join(c for c in s if not unicodedata.combining(c)).lower().strip()


def je_na(r) -> bool:
    return r['attrs'].get('Tip') in (None, '', 'N/A')


def ljudska_oznaka(r) -> str:
    """Komentar koji je netko NAPISAO, a ne prepisani tekst izvoda. '' = nema."""
    c = str(r.get('comment') or '').strip()
    io = str(r['attrs'].get('Izvod opis') or '').strip()
    if not c or len(c) > MAX_OZNAKA:
        return ''
    if io and (fold(io).startswith(fold(c)[:12]) or fold(c) in fold(io)):
        return ''        # prepisan izvod (`AMAZON.DE`, `Restoran Lanterna … [kartica: SAŠA]`)
    return fold(RATA_RE.sub(' ', c)).strip(' .,-')


def jednoglasno(zapisi, minimum):
    """zapisi = [(tip, podtip), ...] samo odlučeni. → (tip, podtip, n, ukupno) ili None."""
    if len(zapisi) < minimum:
        return None
    (par, n), = Counter(zapisi).most_common(1)
    return (*par, n, len(zapisi)) if n / len(zapisi) >= MIN_UDIO else None


class Rjecnik:
    """Indeksi po opisu i trgovcu nad SVIM odlučenim retcima Aree; filtar po datumu na upitu."""

    def __init__(self, redci):
        self.po_opisu = defaultdict(list)
        self.po_trgovcu = defaultdict(list)
        self.po_trgovcu_iznosu = defaultdict(list)
        for r in redci:
            a = r['attrs']
            if je_na(r):
                continue            # N/A je izostanak odluke, ne glas
            z = (r['event_date'], (a['Tip'], a.get('Podtip') or ''))
            o = ljudska_oznaka(r)
            if o:
                self.po_opisu[o].append(z)
            if a.get('Izvor') in KARTICE:
                k = kljuc_kartice(a.get('Izvod opis') or r.get('comment') or '')
                if k and not POSREDNIK.match(k):
                    self.po_trgovcu[k].append(z)
                    self.po_trgovcu_iznosu[(k, abs(net(a)))].append(z)

    @staticmethod
    def _uzmi(lista, prije):
        return [p for d, p in lista if prije is None or d < prije]

    def nadji(self, r, prije=None):
        a = r['attrs']
        o = ljudska_oznaka(r)
        if o:
            g = jednoglasno(self._uzmi(self.po_opisu.get(o, []), prije), MIN_PO_IMENU)
            if g:
                return g, f'opis „{o}"'
        if a.get('Izvor') in KARTICE:
            k = kljuc_kartice(a.get('Izvod opis') or r.get('comment') or '')
            if not k:
                return None, 'nema teksta izvoda'
            if POSREDNIK.match(k):
                return None, f'posrednik „{k}" — izvod ne kaže što je kupljeno'
            par, dokaz = rucno(k, abs(net(a)))
            if par:
                return (*par, 1, 1), dokaz
            g = jednoglasno(self._uzmi(self.po_trgovcu.get(k, []), prije), MIN_PO_IMENU)
            if g:
                return g, f'trgovac „{k}"'
            g = jednoglasno(self._uzmi(self.po_trgovcu_iznosu.get((k, abs(net(a))), []), prije), MIN_PO_IMENU)
            if g:
                return g, f'trgovac „{k}" + iznos {abs(net(a)):.2f}'
            pp = self._uzmi(self.po_trgovcu.get(k, []), prije)
            if not pp:
                return None, f'trgovac „{k}" bez presedana'
            if len(pp) < MIN_PO_IMENU:
                return None, f'trgovac „{k}": samo {len(pp)} presedan(a), treba {MIN_PO_IMENU}'
            return None, f'trgovac „{k}" nije jednoglasan: ' + ', '.join(
                f'{t} / {p} {n}' for (t, p), n in Counter(pp).most_common(2))
        return None, None   # račun — rješava Presedani


def predlozi(redci, ciljevi, rj, tax, provjera=False):
    """→ [(redak, (tip, podtip) | None, dokaz)]"""
    # Presedani su po RAČUNU i izvoru `Racun` (presedani.py: „povijest JEDNOG računa")
    racunski = defaultdict(list)
    for r in redci:
        if r['attrs'].get('Izvor') == 'Racun':
            racunski[r['attrs'].get('Racun')].append(r)
    pres_cache = {}
    out = []
    for r in ciljevi:
        a = r['attrs']
        prije = r['event_date'] if provjera else None
        par, dokaz = rj.nadji(r, prije)
        if par is None and dokaz is None:
            if not a.get('Izvod opis'):
                out.append((r, None, 'nema teksta izvoda ni oznake'))
                continue
            ck = (a.get('Racun'), prije)
            if ck not in pres_cache:
                pres_cache[ck] = Presedani(racunski[a.get('Racun')], prije=prije)
            iznos = net(a)
            smjer = 'Uplata' if iznos < 0 else 'Isplata'
            p = pres_cache[ck].nadji(abs(iznos), a.get('Izvod opis'), smjer)
            if p:
                par = (p['tip'], p['podtip'], *map(int, re.findall(r'(\d+)/(\d+)', p['dokaz'])[0]))
                dokaz = 'račun: ' + p['dokaz'].split(',')[0]
            else:
                dokaz = 'račun: nema jednoglasnog presedana'
        if par:
            tip, podtip, n, uk = par
            if tip not in tax or (podtip and podtip not in tax[tip]):
                out.append((r, None, f'{tip} / {podtip} nije u popisu opcija — ne predlažem'))
                continue
            out.append((r, (tip, podtip), f'{dokaz}: {tip} / {podtip}' + ('' if dokaz.startswith('ručni') else f' {n}/{uk}')))
        else:
            out.append((r, None, dokaz))
    return out


def preuzmi(rez, fileovi, tax):
    """Tip/Podtip koje je čovjek već upisao u ranije razvrstan file, po `event_id`. S164.

    ZAŠTO: razvrstava se jednom (npr. na TEST-u, za probu), a uvozi na PROD. TEST file se na
    PROD NE SMIJE uvesti: kol. G nosi TEST vlasnika, a PROD retke je upisala Koka ⇒ uvoz stane
    (S149) — ili bi, bez te brane, napravio duplikate (S148). TEST je kopija PROD-a s ISTIM
    event ID-evima (`prod_to_test.py`), pa se odluka prenese po ID-u u svjež PROD file, s
    autorom iz PROD baze. Ručna odluka pobjeđuje prijedlog alata.
    """
    from openpyxl import load_workbook
    rucno = {}
    for f in fileovi:
        ws = load_workbook(f, read_only=True)['Events']
        redovi = list(ws.iter_rows(values_only=True))
        hdr = next(i for i, r in enumerate(redovi) if r and r[0] == 'event_id')
        h = list(redovi[hdr])
        it, ip = h.index('Tip'), h.index('Podtip')
        for r in redovi[hdr + 1:]:
            if r and r[0] and r[it] not in (None, '', 'N/A'):
                rucno[str(r[0])] = (str(r[it]), str(r[ip] or ''))
        print(f'Preuzimam iz {Path(f).name}: {sum(1 for _ in rucno)} razvrstanih redaka (ukupno)')
    out, n, krivo = [], 0, []
    for r, p, dokaz in rez:
        par = rucno.get(r['id'])
        if par and (par[0] not in tax or (par[1] and par[1] not in tax[par[0]])):
            krivo.append(f"{r['event_date']} {par[0]} / {par[1]}")
            par = None
        if par:
            n += 1
            out.append((r, par, 'PREUZETO (ručno razvrstano)' if par != p else dokaz))
        else:
            out.append((r, p, dokaz))
    print(f'  primijenjeno na {n} N/A redaka ove baze')
    if krivo:
        print(f'  ✗ {len(krivo)} par(ova) nije u popisu opcija — NISU preuzeti: ' + '; '.join(krivo))
    return out


def main():
    env = target()
    url, key = load_env(env)
    od_, do_ = prozor()
    print(f'[{env.upper()}] {url}\nProzor: {od_} … {do_}\n')
    redci = load_db(url, key)
    tax = taksonomija(url, key)
    rj = Rjecnik(redci)
    u_prozoru = [r for r in redci if od_ <= r['event_date'] <= do_]

    if '--provjera' in sys.argv:
        ciljevi = [r for r in u_prozoru if not je_na(r) and r['attrs'].get('Tip') != 'Transfer']
        rez = predlozi(redci, ciljevi, rj, tax, provjera=True)
        dano = [(r, p) for r, p, _ in rez if p]
        tocno = sum(1 for r, p in dano if p == (r['attrs']['Tip'], r['attrs'].get('Podtip') or ''))
        print(f'Provjera unatrag: {len(ciljevi)} razvrstanih redaka (bez Transfera), povijest starija od retka')
        print(f'  prijedlog dan: {len(dano)} ({len(dano) / max(len(ciljevi), 1):.0%})')
        print(f'  od toga točan: {tocno} ({tocno / max(len(dano), 1):.1%})')
        krivi = Counter(((r['attrs']['Tip'], r['attrs'].get('Podtip')), p) for r, p in dano
                        if p != (r['attrs']['Tip'], r['attrs'].get('Podtip') or ''))
        for (stvarno, predlozeno), n in krivi.most_common(10):
            print(f'    {n:>3}×  stvarno {stvarno[0]} / {stvarno[1]}  ←  predloženo {predlozeno[0]} / {predlozeno[1]}')
        return

    ciljevi = sorted([r for r in u_prozoru if je_na(r)], key=lambda r: r['event_date'])
    rez = predlozi(redci, ciljevi, rj, tax)
    if '--preuzmi' in sys.argv:
        rez = preuzmi(rez, [a for a in sys.argv[sys.argv.index('--preuzmi') + 1:]
                            if not a.startswith('--')], tax)
    po_racunu = defaultdict(list)
    for x in rez:
        po_racunu[x[0]['attrs'].get('Racun') or '(bez računa)'].append(x)

    uk_p = sum(1 for _, p, _ in rez if p)
    print(f'N/A redaka: {len(rez)} · s prijedlogom {uk_p} · ostaje ručno {len(rez) - uk_p}\n')
    for racun, xs in sorted(po_racunu.items()):
        s = sum(1 for _, p, _ in xs if p)
        print(f'── {racun}: {len(xs)} redaka, prijedlog {s}')
        for r, p, dokaz in xs:
            iznos = net(r['attrs'])
            opis = (r.get('comment') or r['attrs'].get('Izvod opis') or '')[:38]
            print(f"  {r['event_date']} {iznos:>9,.2f} {r['attrs'].get('Izvor', ''):<10} {opis:<38} "
                  f"{'→ ' + p[0] + ' / ' + p[1] if p else '·'}   [{dokaz}]")
        print()

    if '--file' not in sys.argv:
        print('Ispis. Za Excel po računu: --file')
        return

    # e-mailovi autora iz baze u koju gađamo (TEST kopija ima svoga vlasnika)
    for pr in rest(url, key, 'profiles?select=id,email'):
        if pr.get('email'):
            visa_popravak.EMAIL.setdefault(pr['id'], pr['email'])
    for racun, xs in sorted(po_racunu.items()):
        redci_xl = []
        for r, p, dokaz in xs:
            a = dict(r['attrs'])
            if p:
                a['Tip'], a['Podtip'] = p
            redci_xl.append((r, a, r.get('comment'), False,
                             (f'PRIJEDLOG {p[0]} / {p[1]} — ' if p else 'RUČNO — ') + str(dokaz)))
        ime = 'na_razvrstavanje_' + re.sub(r'[^A-Za-z0-9]+', '_', fold(racun)).strip('_')
        visa_popravak.pisi(redci_xl, {}, ime, tax=tax, izgled=izgled_pregleda)
    print('\nUvoz: vlasnica Aree (Koka, ili Saša pod njezinim računom), tuđi retci „fix as owner".'
          '\nOčekivani pregled: 0 New · Modify = broj redaka koje je pregled ostavio s Tipom · 0 Delete.')


if __name__ == '__main__':
    main()
