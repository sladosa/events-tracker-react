# -*- coding: utf-8 -*-
"""
rate_alat.py — rate: higijena oznaka (prolaz A) + generiranje preostalih (prolaz B). S130.

ZASTO POSTOJI
    Baza drzi 636 redaka s `Rate? = true`, ali **u modelu ne postoji identifikator
    plana**. Postoje `Rate?`, `Broj rata`, `Rata br` -- nista sto kaze "ovi retci su
    jedna kupovina". Plan se zato rekonstruira iz `Izvod opis`a, koji za 433 retka
    nosi bankin tekst oblika `KONZUM P-0277 RATA 9/12`.

AUTORITET JE `Izvod opis`, NIKAD KOMENTAR
    Izmjereno: Kokin `Rata br` se s bankinim brojem slaze u 352 slucaja, razlikuje u
    15 (9x za -1, 6x za +1), a u 66 ga uopce nema. Bankin tekst je jedini izvor koji
    nosi i broj rate i ukupan broj i trgovca.

    /!\ PRAVILO PO KOMENTARU JE ODBACENO -- MJERENJEM, NE OSJECAJEM.
        Uzorak `n/N` u komentaru hvata 49 redaka, a stvarne rate su tri:
            HLK clanarina 7-12/23   -> razdoblje (srpanj-prosinac), NE rata
            HLK 03/23 7,96          -> mjesec
            HLK 4,5,6/25            -> tri mjeseca
        Isti razred kao "pretraga po kljucnoj rijeci prekomjerno hvata" (CLAUDE.md).
        Zato retci bez `Izvod opis`a idu kroz IZMJEREN POPIS (`RUCNO`), ne kroz uzorak.

/!\ PRVA RATA NOSI OSTATAK ZAOKRUZIVANJA
    Izmjereno na 62 plana s >=3 rate: 25 ima sve iznose jednake, **23 ima drukciju
    samo PRVU** (`DECATHLON 1: 7,82`, pa `2..12: 7,79`), 14 je sudar kljuca.
    Zato se za nove rate uzima iznos **ZADNJE** rate, nikad prve.

/!\ MC dospijece je pravilno, VISA NIJE
    Izmjereno: MC `Datum naplate` je 11. u mjesecu u **276/276** slucajeva.
    Visa: 5. (113), 4. (71), 6. (29), 7. (17), 12. (12), 8. (12) -- dakle nema pravila.
    Zato se generiraju SAMO MC rate; Visa se prijavljuje i ceka izvod.

/!\ `event_date` = DAN KUPNJE (Sasina odluka S130)
    Sve rate jedne kupovine dijele dan kupnje; razlikuje ih `Datum naplate`.
    Kolizija se izbjegava pomakom `session_start`a u pojasu 14:00+, jer
    `useActivities` grupira po (user, kategorija, session_start) -- dva retka iste
    minute postaju JEDAN redak liste.

/!\ ZASTAREO PLAN SE NE GENERIRA
    Plan cija bi sljedeca rata dospjela u PROSLOSTI nije otvoren nego nezabiljezen
    (`HARVEY NORMAN`, zadnja 28.09.2024.). Generirati ga znaci proizvesti `Planiran`
    redak s proslim dospijecem, sto delta sheet odmah prijavi kao gresku. Prijavljuje
    se, ne generira.

Pokretanje:
    python rate_alat.py                      # dry run, samo ispis
    python rate_alat.py --file rate          # + rate_A.xlsx i rate_B.xlsx
    python rate_alat.py --only a --file rate # samo prolaz A

/!\ Alat NE PISE u bazu. Proizvodi xlsx koji se uvozi kroz aplikaciju, pod Kokinim
    racunom (ona je vlasnica Aree).
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from uskladi_izvod import load_db, load_env, net  # noqa: E402

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.exit('pip install openpyxl')

AREA = 'Financije_all'
CAT_PATH = 'Transakcija'          # BEZ imena aree (v. excel_import_template.py)
KOKA_EMAIL = 'dubravka.pavic-sladoljev@dps-perceptum.com'

RATA_PAT = re.compile(r'RATA\s*(\d+)\s*/\s*(\d+)', re.I)

# Atributi koje pisemo. Redoslijed odredjuje Excel stupac (I, J, K, ...).
ATTRS = [
    ('Racun',         'text'),
    ('Izvor',         'text'),
    ('Smjer',         'text'),
    ('Uplata',        'number'),
    ('Isplata',       'number'),
    ('Tip',           'text'),
    ('Podtip',        'text'),
    ('Izvod opis',    'text'),
    ('Rate?',         'boolean'),
    ('Broj rata',     'number'),
    ('Rata br',       'number'),
    ('Datum naplate', 'datetime'),
    ('Status',        'text'),
]

# ---------------------------------------------------------------------------
# IZMJEREN popis retka koji nose ratu, a `Izvod opis` im je ne moze potvrditi.
# /!\ POPIS, NE PRAVILO -- v. zaglavlje. Svaki je provjeren pojedinacno.
#     (komentar koji ga identificira, Broj rata, Rata br)
# ---------------------------------------------------------------------------
RUCNO = [
    ('Anja 84/96',  96, 84),
    ('Anja 85/96',  96, 85),
    ('Konzum 1/6',   6,  1),
]


# ------------------------------------------------------------------ helpers --
def plus_months(d: date, n: int) -> date:
    """Isti dan u mjesecu, n mjeseci naprijed (dan se ne pomice -- MC je uvijek 11.)."""
    m = d.month - 1 + n
    return date(d.year + m // 12, m % 12 + 1, d.day)


def parse_due(v) -> date | None:
    if not v:
        return None
    return datetime.fromisoformat(str(v)[:10]).date()


def hr(d: date) -> str:
    return d.strftime('%d.%m.%Y.')


# ------------------------------------------------------------------ analiza --
def plans(db):
    """Planovi rekonstruirani iz `Izvod opis`a. Kljuc = (trgovac, N)."""
    g = defaultdict(list)
    for r in db:
        m = RATA_PAT.search(str(r['attrs'].get('Izvod opis') or ''))
        if m:
            trg = RATA_PAT.sub('', str(r['attrs']['Izvod opis'])).strip()
            g[(trg, int(m.group(2)))].append((int(m.group(1)), r))
    return g


def prolaz_a(db):
    """Retci kojima fali ili se ne slaze oznaka rate / status."""
    fix = defaultdict(dict)          # event_id -> {polje: nova vrijednost}
    zasto = {}
    for r in db:
        a = r['attrs']
        m = RATA_PAT.search(str(a.get('Izvod opis') or ''))
        if m:
            n, N = int(m.group(1)), int(m.group(2))
            if a.get('Rate?') is not True:
                fix[r['id']]['Rate?'] = True
            if a.get('Broj rata') != N:
                fix[r['id']]['Broj rata'] = N
            b = a.get('Rata br')
            if b is None or int(b) != n:
                fix[r['id']]['Rata br'] = n
            if r['id'] in fix:
                zasto[r['id']] = 'izvod: ' + str(a['Izvod opis'])[:40]
        if not a.get('Status'):
            fix[r['id']]['Status'] = 'Planiran'
            zasto.setdefault(r['id'], 'prazan Status')

    for kom, N, n in RUCNO:
        for r in db:
            if str(r['comment'] or '').strip() != kom:
                continue
            a = r['attrs']
            if a.get('Rate?') is not True:
                fix[r['id']]['Rate?'] = True
            if a.get('Broj rata') != N:
                fix[r['id']]['Broj rata'] = N
            if a.get('Rata br') != n:
                fix[r['id']]['Rata br'] = n
            if r['id'] in fix:
                zasto.setdefault(r['id'], 'rucni popis')
    return {k: v for k, v in fix.items() if v}, zasto


def prolaz_b(db, danas: date):
    """Preostale rate otvorenih planova."""
    nove, visa, zastarjeli, uspor = [], [], [], []
    for (trg, N), v in plans(db).items():
        br = [x[0] for x in v]
        if len(br) != len(set(br)):
            uspor.append((trg, N, sorted(set(round(abs(net(x[1]['attrs'])), 2) for x in v))))
            continue
        mx = max(br)
        if mx >= N:
            continue
        zadnja = [x[1] for x in v if x[0] == mx][0]
        a = zadnja['attrs']
        due = parse_due(a.get('Datum naplate'))
        prva = min(v)[1]
        red = {
            'trgovac': trg, 'N': N, 'od': mx + 1, 'iznos': round(abs(net(a)), 2),
            'event_date': prva['event_date'], 'izvor': a.get('Izvor'),
            'racun': a.get('Racun'), 'tip': a.get('Tip'), 'podtip': a.get('Podtip'),
            'due': due, 'comment': str(zadnja['comment'] or ''),
        }
        if a.get('Izvor') != 'Mastercard':
            visa.append(red)
            continue
        if not due or plus_months(due, 1) < danas:
            zastarjeli.append(red)
            continue
        nove.append(red)
    return nove, visa, zastarjeli, uspor


def slobodne_minute(db, dan: str, koliko: int) -> list[str]:
    """/!\\ Kolizija `session_start`a spaja retke u JEDAN redak liste (useActivities)."""
    zauzete = {str(r.get('session_start') or '')[11:16] for r in db if r['event_date'] == dan}
    out = []
    for h in range(14, 24):
        for mi in range(60):
            s = '%02d:%02d' % (h, mi)
            if s not in zauzete:
                out.append(s)
                if len(out) == koliko:
                    return out
    return out


# ------------------------------------------------------------------- zapis --
def tekst(ws, r, c, v):
    """/!\\ openpyxl string koji pocinje s `=` sprema kao FORMULU -> file se ne otvori."""
    cell = ws.cell(r, c)
    if isinstance(v, str) and v[:1] in ('=', '+', '-', '@'):
        cell.value = v
        cell.data_type = 's'
    else:
        cell.value = v
    return cell


def col_letter(n):
    s = ''
    while n > 0:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


def napisi(path: Path, redci: list[dict]):
    """`redci`: {event_id?, date, time, comment, attrs:{ime: vrijednost}}"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Events'
    purple = PatternFill('solid', fgColor='7030A0')
    blue = PatternFill('solid', fgColor='4472C4')
    white = Font(bold=True, color='FFFFFF')

    ws.cell(1, 1, 'ATTRIBUTE LEGEND:').font = Font(bold=True, size=12)
    for i, h in enumerate(['Col', 'Area', 'Category_Path', 'Attribute', 'Type', 'Unit'], start=1):
        c = ws.cell(2, i, h)
        c.fill, c.font, c.alignment = purple, white, Alignment(horizontal='center')
    for i, (name, typ) in enumerate(ATTRS):
        r = 3 + i
        ws.cell(r, 1, col_letter(9 + i))
        ws.cell(r, 2, AREA)
        ws.cell(r, 3, CAT_PATH)
        ws.cell(r, 4, name)
        ws.cell(r, 5, typ)

    top = 3 + len(ATTRS) + 1
    ws.cell(top, 1, 'EVENT DATA:').font = Font(bold=True, size=12)
    hdr = top + 1
    heads = ['event_id', 'Area', 'Category_Path', 'event_date', 'session_start',
             'created_at', 'User', 'leaf comment'] + [n for n, _ in ATTRS]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(hdr, i, h)
        c.fill, c.font = blue, white

    for j, row in enumerate(redci):
        r = hdr + 1 + j
        tekst(ws, r, 1, row.get('event_id') or None)
        tekst(ws, r, 2, AREA)
        tekst(ws, r, 3, CAT_PATH)
        tekst(ws, r, 4, row['date'])
        # /!\ MORA biti TEKST "HH:MM" -- prava Excel time vrijednost daje puni ISO,
        #     `parseTimeStr` vrati null i SVI redci padnu na 09:00.
        c = ws.cell(r, 5, row['time'])
        c.number_format = '@'
        tekst(ws, r, 7, KOKA_EMAIL)
        tekst(ws, r, 8, row.get('comment') or None)
        for i, (name, typ) in enumerate(ATTRS):
            v = row['attrs'].get(name)
            if v is None:
                continue
            col = 9 + i
            if typ == 'boolean':
                # /!\ sve osim doslovnog 'true' sprema se kao FALSE, bez poruke
                tekst(ws, r, col, 'true' if v else 'false')
            elif typ == 'datetime':
                # /!\ datumska celija se sidri u PODNE UTC (v. excelDatetime.ts)
                ws.cell(r, col, datetime(v.year, v.month, v.day, 12, 0))
            else:
                tekst(ws, r, col, v)

    for i, w in enumerate([12, 14, 16, 12, 12, 10, 34, 30], start=1):
        ws.column_dimensions[col_letter(i)].width = w
    for i in range(len(ATTRS)):
        ws.column_dimensions[col_letter(9 + i)].width = 16
    wb.save(path)


# -------------------------------------------------------------------- main --
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--env', default='prod', choices=['prod', 'test'])
    ap.add_argument('--file', help='prefiks; pise <prefiks>_A.xlsx i <prefiks>_B.xlsx')
    ap.add_argument('--only', choices=['a', 'b'], help='samo jedan prolaz')
    ap.add_argument('--danas', help='YYYY-MM-DD (za test)')
    args = ap.parse_args()

    danas = datetime.fromisoformat(args.danas).date() if args.danas else date.today()
    url, key = load_env(args.env)
    db = load_db(url, key)
    by_id = {r['id']: r for r in db}
    print('=' * 96)
    print('RATE — higijena i generiranje   [%s]   %d redaka u bazi   danas %s'
          % (args.env.upper(), len(db), hr(danas)))
    print('=' * 96)

    a_redci, b_redci = [], []

    # ---------------------------------------------------------------- A --
    if args.only != 'b':
        fix, zasto = prolaz_a(db)
        print('\nPROLAZ A — higijena oznaka   [%d redaka]' % len(fix))
        print('-' * 96)
        for eid, polja in sorted(fix.items(), key=lambda x: by_id[x[0]]['event_date']):
            r = by_id[eid]
            print('  %s  %-26s %8.2f   %s'
                  % (r['event_date'], str(r['comment'] or '')[:26], net(r['attrs']),
                     ', '.join('%s=%s' % (k, v) for k, v in polja.items())))
            print('        %s' % zasto.get(eid, ''))
            a_redci.append({
                'event_id': eid, 'date': r['event_date'],
                'time': str(r.get('session_start') or '')[11:16] or '14:00',
                'comment': None, 'attrs': polja,
            })

    # ---------------------------------------------------------------- B --
    if args.only != 'a':
        nove, visa, zastarjeli, uspor = prolaz_b(db, danas)
        uk = sum(p['N'] - p['od'] + 1 for p in nove)
        print('\nPROLAZ B — nove rate   [%d planova, %d rata]' % (len(nove), uk))
        print('-' * 96)
        for p in sorted(nove, key=lambda x: -(x['N'] - x['od'] + 1)):
            slob = slobodne_minute(db, p['event_date'], p['N'] - p['od'] + 1)
            print('  %-26s %6.2f  rate %d..%d  dan kupnje %s  minute %s'
                  % (p['trgovac'][:26], p['iznos'], p['od'], p['N'], p['event_date'],
                     ', '.join(slob)))
            for k, n in enumerate(range(p['od'], p['N'] + 1)):
                due = plus_months(p['due'], n - p['od'] + 1)
                print('        rata %2d/%d   dospijece %s' % (n, p['N'], hr(due)))
                b_redci.append({
                    'date': p['event_date'], 'time': slob[k] if k < len(slob) else '23:%02d' % k,
                    'comment': re.sub(r'\d+/\d+\s*$', '', p['comment']).strip()
                               + ' %d/%d' % (n, p['N']),
                    'attrs': {
                        'Racun': p['racun'], 'Izvor': p['izvor'], 'Smjer': 'Isplata',
                        'Isplata': p['iznos'], 'Tip': p['tip'], 'Podtip': p['podtip'],
                        'Rate?': True, 'Broj rata': p['N'], 'Rata br': n,
                        'Datum naplate': due, 'Status': 'Planiran',
                    },
                })

        print('\n  NE GENERIRA SE — Visa (dospijece nije pravilno)   [%d planova, %d rata]'
              % (len(visa), sum(p['N'] - p['od'] + 1 for p in visa)))
        for p in sorted(visa, key=lambda x: -(x['N'] - x['od'] + 1))[:8]:
            print('     %-34s %7.2f  fali %d  (zadnja %d/%d)'
                  % (p['trgovac'][:34], p['iznos'], p['N'] - p['od'] + 1, p['od'] - 1, p['N']))

        print('\n  NE GENERIRA SE — zastarjeli (sljedeca rata bi dospjela u PROSLOSTI)   [%d]'
              % len(zastarjeli))
        for p in zastarjeli:
            print('     %-34s %7.2f  zadnja %d/%d, dospijece %s'
                  % (p['trgovac'][:34], p['iznos'], p['od'] - 1, p['N'],
                     hr(p['due']) if p['due'] else '?'))

        print('\n  NE GENERIRA SE — usporedni planovi (isti trgovac+N, razdvaja ih iznos)   [%d]'
              % len(uspor))
        for trg, N, izn in uspor:
            print('     %-34s N=%-3d iznosi %s' % (trg[:34], N, izn))

    # ------------------------------------------------------------- zapis --
    print('\n' + '=' * 96)
    if not args.file:
        print('DRY RUN — nista nije napisano. Za xlsx: --file rate')
        return
    out = Path(args.file)
    if a_redci:
        p = out.with_name(out.name + '_A.xlsx')
        napisi(p, a_redci)
        print('napisano: %s   (%d redaka, svi UPDATE)' % (p, len(a_redci)))
    if b_redci:
        p = out.with_name(out.name + '_B.xlsx')
        napisi(p, b_redci)
        print('napisano: %s   (%d redaka, svi CREATE)' % (p, len(b_redci)))
    print('\nUvoz ide kroz aplikaciju, pod Kokinim racunom. Prvo A, provjeri, pa B.')


if __name__ == '__main__':
    main()
