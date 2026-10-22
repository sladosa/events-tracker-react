# -*- coding: utf-8 -*-
"""Gdje Python alati jos citiraju vrijednost koju sheet `Preuredjivanje` mijenja (F1, S179).

Planer u appu prepise SVE sto zivi u bazi (retci, izbornik, pločice, bucketi,
koske, kolone, shortcuti, sidra) -- ali preglednik ne vidi Python rjecnike
(`pravila_trgovca.PRAVILA`, `visa_uvoz_izvoda.KLASA`, `set_breakdown.py`...).
Oni bi i dalje predlagali STARI naziv, bez greske (zamka S167). Ovaj alat ih
nabroji prije primjene; ispravljaju se istog dana (PREUREDJIVANJE_SPEC §16.7).

Pokretanje iz korijena projekta (suhi prolaz, nista ne pise):
    Tools\\run.bat Tools\\preuredjivanje_grep.py <structure.xlsx>
    Tools\\run.bat Tools\\preuredjivanje_grep.py <structure.xlsx> --sve   # i slabi pogoci

STO JE „STARA VRIJEDNOST"
    Za svaki korak i svaki atribut iz „postavi" koji ima i „kad" s drugom
    vrijednoscu: `kad` vrijednost je stara. Ostali „kad" uvjeti istog koraka su
    KONTEKST: korak 3 (kad tip=Kuca, podtip=Osiguranje -> tip=Osiguranje) ne gasi
    `Kuca` -- gasi PAR. Zato:
      JAK  pogodak = redak sadrzi staru vrijednost I sve vrijednosti konteksta
           (ili korak nema konteksta: `Visa` -> `VISA`);
      SLAB pogodak = samo stara vrijednost -- broji se po fileu, ispis uz --sve.
    Trazi se DOSLOVAN string literal ('X' ili "X"), ne podniz: `Sasa` ne smije
    pogoditi `Sasa_tekuci`.

/!\\ `fix_*.py` i `Obsolete/` se preskacu: to je ZAPIS sto je uradjeno nad
    podacima (CLAUDE.md, „87 Python alata"), ne alat koji ce sutra predloziti
    stari naziv. Ispravljati povijest bi bilo krivotvorenje zapisa.
/!\\ Koraci samo s `event_id` (bez „kad" za taj atribut) nemaju staru vrijednost
    -- pojedinacne odluke ne gase naziv.

Izlaz: 0 = nema jakih pogodaka, 1 = ima (popis je posao za isti dan).
"""
import argparse
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

ROOT = Path(__file__).resolve().parents[2]
SCAN = ROOT / 'data-prep_tools'
SKIP_PARTS = {'Obsolete', '__pycache__', 'venv', 'site-packages', '_arhiva'}
EMPTY = '(prazno)'
WHEN, SET = 'kad:', 'postavi:'


def _norm(s):
    s = unicodedata.normalize('NFD', s.lower())
    return ''.join(c for c in s if not unicodedata.combining(c)).replace('đ', 'd')


def read_steps(path):
    """Sheet `Preuređivanje` -> [(korak, {slug: kad}, {slug: postavi})]. Isto kao `restructure/sheet.ts`."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = next((w for w in wb.worksheets if _norm(w.title) in ('preuredivanje', 'preuredjivanje')), None)
    if ws is None:
        sys.exit(f'Nema sheeta „Preuređivanje" u {path}.')
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [str(h).strip() if h is not None else '' for h in rows[0]]
    cols = []
    for i, h in enumerate(head):
        low = h.lower()
        if low == 'korak':
            cols.append(('n', i, None))
        elif low.startswith(WHEN):
            cols.append(('when', i, h[len(WHEN):].strip()))
        elif low.startswith(SET):
            cols.append(('set', i, h[len(SET):].strip()))
    steps = []
    for r in rows[1:]:
        if r and isinstance(r[0], str) and r[0].startswith('HELP'):
            break
        if not r or all(c is None or str(c).strip() == '' for c in r):
            continue
        n, when, setv = None, {}, {}
        for kind, i, slug in cols:
            v = r[i] if i < len(r) else None
            if v is None or str(v).strip() == '':
                continue
            if kind == 'n':
                n = v
                continue
            v = None if str(v).strip() == EMPTY else str(v)
            (when if kind == 'when' else setv)[slug] = v
        steps.append((n, when, setv))
    return steps


def old_values(steps):
    """[(korak, slug, stara, {kontekst slug: vrijednost}, nova)]"""
    out = []
    for n, when, setv in steps:
        for slug, new in setv.items():
            old = when.get(slug)
            if old is None or old == new:
                continue
            ctx = {s: v for s, v in when.items() if s != slug and v is not None}
            out.append((n, slug, old, ctx, new))
    return out


def lit(v):
    return re.compile(r'''(['"])''' + re.escape(v) + r'\1')


def files():
    for p in sorted(SCAN.rglob('*.py')):
        if SKIP_PARTS & set(p.parts) or p.name.startswith('fix_') or p.name == Path(__file__).name:
            continue
        yield p


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('xlsx')
    ap.add_argument('--sve', action='store_true', help='ispiši i slabe pogotke (samo vrijednost, bez konteksta)')
    a = ap.parse_args()

    olds = old_values(read_steps(a.xlsx))
    if not olds:
        print('Nijedan korak ne gasi postojeću vrijednost (nema „kad X → postavi Y" za isti atribut).')
        return 0

    texts = {p: p.read_text(encoding='utf-8', errors='replace').splitlines() for p in files()}
    strong_total = 0
    for n, slug, old, ctx, new in olds:
        ctx_txt = ', '.join(f'{s}={v}' for s, v in ctx.items())
        print(f'\n── korak {n}: {slug} „{old}" → „{new}"' + (f'   (uz {ctx_txt})' if ctx else ''))
        pat = lit(old)
        ctx_pats = [lit(v) for v in ctx.values()]
        strong, weak = [], defaultdict(list)
        for p, lines in texts.items():
            for i, line in enumerate(lines, 1):
                if not pat.search(line):
                    continue
                if all(c.search(line) for c in ctx_pats):
                    strong.append((p, i, line.strip()))
                else:
                    weak[p].append((i, line.strip()))
        strong_total += len(strong)
        for p, i, line in strong:
            print(f'  JAK   {p.relative_to(ROOT)}:{i}   {line[:140]}')
        if not strong:
            print('  (nema jakih pogodaka)')
        if weak:
            nweak = sum(len(v) for v in weak.values())
            print(f'  slabi: {nweak} redak/a u {len(weak)} fileu/ova — samo „{old}", bez {ctx_txt or "konteksta"}'
                  + ('' if a.sve else ' (--sve za popis)'))
            if a.sve:
                for p, hits in weak.items():
                    for i, line in hits:
                        print(f'  slab  {p.relative_to(ROOT)}:{i}   {line[:140]}')

    print(f'\nJakih pogodaka: {strong_total}. Pregledano {len(texts)} Python fileova u data-prep_tools/'
          f' (bez fix_*.py i Obsolete/).')
    return 1 if strong_total else 0


if __name__ == '__main__':
    sys.exit(main())
